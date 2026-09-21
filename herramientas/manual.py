#!/usr/bin/env python3
"""
Genera el manual de usuario de la carta: manual/Manual carta Boró Café.pdf

Uso (necesita reportlab y openpyxl; fuentes del sistema de macOS):
    python3 herramientas/manual.py

La miniatura de la carta se dibuja en el PDF (clase Maqueta): no son capturas, así no envejece.
La muestra de la planilla se lee de datos/carta.xlsx, así que refleja la planilla vigente.
"""
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (BaseDocTemplate, CondPageBreak, Flowable, Frame, Image, KeepTogether,
                                NextPageTemplate, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle)

RAIZ = Path(__file__).resolve().parent.parent
MANUAL = RAIZ / "manual"
SALIDA = MANUAL / "Manual carta Boró Café.pdf"

URL_CARTA = "https://borocafe.github.io/carta/"
URL_DESCARGA = "https://github.com/borocafe/carta/raw/main/datos/carta.xlsx"
URL_SUBIDA = "https://github.com/borocafe/carta/upload/main/datos"
URL_ACTIONS = "https://github.com/borocafe/carta/actions"
URL_HISTORIAL = "https://github.com/borocafe/carta/commits/main/datos"
URL_QR = "https://borocafe.github.io/carta/ir/"
URL_TARJETA = "https://github.com/borocafe/carta/raw/main/qr/tarjeta-mesa.pdf"

OLIVA = colors.HexColor("#4B5324")
MARCA = colors.HexColor("#5F672A")
CREMA = colors.HexColor("#F4E8D0")
CREMA_SUAVE = colors.HexColor("#FAF5EA")
NOTA_FONDO = colors.HexColor("#F6F0E2")
TINTA = colors.HexColor("#3B3026")
TINTA_SUAVE = colors.HexColor("#6D5B48")
LINEA = colors.HexColor("#D9C7A7")
PAPEL_HONDO = colors.HexColor("#EDDFC2")
ALERTA = colors.HexColor("#9A3F24")
ALERTA_FONDO = colors.HexColor("#F7E7DE")

ANCHO_PAGINA, ALTO_PAGINA = A4
MARGEN_X, MARGEN_SUP, MARGEN_INF = 22 * mm, 24 * mm, 20 * mm
ANCHO = ANCHO_PAGINA - 2 * MARGEN_X
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
         "noviembre", "diciembre"]

AVENIR = "/System/Library/Fonts/Avenir Next.ttc"
BASKERVILLE = "/System/Library/Fonts/Supplemental/Baskerville.ttc"


def registrar_fuentes():
    for nombre, ruta, indice in [("Avenir", AVENIR, 7), ("Avenir-Medio", AVENIR, 5), ("Avenir-Demi", AVENIR, 2),
                                 ("Avenir-Italica", AVENIR, 4), ("Avenir-DemiItalica", AVENIR, 3),
                                 ("Basker", BASKERVILLE, 4), ("Basker-Italica", BASKERVILLE, 2)]:
        pdfmetrics.registerFont(TTFont(nombre, ruta, subfontIndex=indice))
    pdfmetrics.registerFontFamily("Avenir", normal="Avenir", bold="Avenir-Demi", italic="Avenir-Italica",
                                  boldItalic="Avenir-DemiItalica")


E = {}


def estilos():
    base = dict(fontName="Avenir", textColor=TINTA)
    E.update(
        cuerpo=ParagraphStyle("cuerpo", fontSize=10, leading=15, **base),
        chico=ParagraphStyle("chico", fontSize=8.5, leading=12.5, fontName="Avenir", textColor=TINTA_SUAVE),
        tabla=ParagraphStyle("tabla", fontSize=9, leading=12.5, **base),
        tabla_nota=ParagraphStyle("tabla_nota", fontSize=9, leading=12.5, fontName="Avenir-Italica", textColor=TINTA_SUAVE),
        tabla_cab=ParagraphStyle("tabla_cab", fontSize=8, leading=10, fontName="Avenir-Demi", textColor=CREMA),
        h1=ParagraphStyle("h1", fontSize=23, leading=28, fontName="Basker", textColor=OLIVA, spaceAfter=4),
        bajada=ParagraphStyle("bajada", fontSize=10.5, leading=15.5, fontName="Avenir", textColor=TINTA_SUAVE, spaceAfter=12),
        h2=ParagraphStyle("h2", fontSize=14, leading=18, fontName="Basker", textColor=OLIVA, spaceBefore=16, spaceAfter=6),
        etiqueta=ParagraphStyle("etiqueta", fontSize=7.5, leading=10, fontName="Avenir-Demi", textColor=MARCA),
        portada_titulo=ParagraphStyle("portada_titulo", fontSize=34, leading=38, fontName="Basker", textColor=OLIVA,
                                      alignment=1),
        portada_bajada=ParagraphStyle("portada_bajada", fontSize=12, leading=18, fontName="Avenir", textColor=TINTA_SUAVE,
                                      alignment=1),
    )


TEXTOS = []  # para revisar que las fuentes tengan todos los caracteres


def P(texto, estilo="cuerpo"):
    TEXTOS.append((texto, E[estilo].fontName))
    return Paragraph(texto, E[estilo])


def enlace(url, texto=None):
    return f'<link href="{url}" color="#5F672A"><b>{texto or url.replace("https://", "")}</b></link>'


def ruta_menu(*partes):
    return " › ".join(f"<b>{p}</b>" for p in partes)


class Numero(Flowable):
    """Círculo oliva con el número de paso."""

    def __init__(self, n, diametro=7 * mm):
        super().__init__()
        self.n, self.d = str(n), diametro
        self.width = self.height = diametro

    def draw(self):
        r = self.d / 2
        self.canv.setFillColor(OLIVA)
        self.canv.circle(r, r, r, stroke=0, fill=1)
        self.canv.setFillColor(CREMA)
        self.canv.setFont("Avenir-Demi", 10)
        self.canv.drawCentredString(r, r - 3.5, self.n)


def pasos(lista, ancho=ANCHO):
    filas = [[Numero(i), [P(f"<b>{titulo}</b>"), P(texto)] if titulo else P(texto)]
             for i, (titulo, texto) in enumerate(lista, start=1)]
    t = Table(filas, colWidths=[11 * mm, ancho - 11 * mm])
    t.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    return t


def recuadro(contenido, fondo=CREMA_SUAVE, borde=LINEA, acento=None, ancho=ANCHO):
    t = Table([[contenido]], colWidths=[ancho])
    estilo = [
        ("BACKGROUND", (0, 0), (-1, -1), fondo),
        ("LEFTPADDING", (0, 0), (-1, -1), 12), ("RIGHTPADDING", (0, 0), (-1, -1), 12),
        ("TOPPADDING", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
    ]
    if acento:
        estilo.append(("LINEBEFORE", (0, 0), (0, -1), 2.5, acento))
    else:
        estilo.append(("BOX", (0, 0), (-1, -1), 0.6, borde))
    t.setStyle(TableStyle(estilo))
    return t


def ojo(texto):
    return recuadro([P("<b>OJO</b>", "etiqueta"), P(texto)], fondo=ALERTA_FONDO, acento=ALERTA)


def tabla(filas, anchos, cabecera=True, cebra=True):
    t = Table(filas, colWidths=anchos, repeatRows=1 if cabecera else 0)
    estilo = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, LINEA),
    ]
    if cabecera:
        estilo += [("BACKGROUND", (0, 0), (-1, 0), OLIVA), ("VALIGN", (0, 0), (-1, 0), "MIDDLE")]
    if cebra:
        inicio = 1 if cabecera else 0
        for i in range(inicio, len(filas)):
            if (i - inicio) % 2 == 1:
                estilo.append(("BACKGROUND", (0, i), (-1, i), CREMA_SUAVE))
    t.setStyle(TableStyle(estilo))
    return t


def imagen(ruta, ancho):
    with PILImage.open(ruta) as im:
        w, h = im.size
    return Image(str(ruta), width=ancho, height=ancho * h / w)


# ---------------------------------------------------------------- contenido

def portada():
    logo = imagen(MANUAL / "img" / "logo.png", 62 * mm)
    enlaces = Table([
        [P("VER LA CARTA", "etiqueta"), P(enlace(URL_CARTA))],
        [P("DESCARGAR LA PLANILLA", "etiqueta"), P(enlace(URL_DESCARGA))],
        [P("SUBIR LA PLANILLA", "etiqueta"), P(enlace(URL_SUBIDA))],
    ], colWidths=[46 * mm, 104 * mm])
    enlaces.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 12), ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
        ("LINEBELOW", (0, 0), (-1, -2), 0.4, LINEA),
    ]))
    hoy = date.today()
    return [
        Spacer(1, 22 * mm), logo, Spacer(1, 14 * mm),
        P("Manual de la carta", "portada_titulo"), Spacer(1, 5 * mm),
        P("Cómo cambiar precios, agregar productos<br/>y publicar la carta digital de Boró Café", "portada_bajada"),
        Spacer(1, 20 * mm),
        Table([[enlaces]], colWidths=[150 * mm], style=[("ALIGN", (0, 0), (-1, -1), "CENTER")]),
        Spacer(1, 16 * mm),
        P(f"Para el equipo de Boró Café · versión {MESES[hoy.month - 1]} {hoy.year}", "portada_bajada"),
        NextPageTemplate("interior"), PageBreak(),
    ]


def resumen():
    return [
        P("La idea en un minuto", "h1"),
        P("Los productos y precios de la carta salen de una planilla Excel. Para cambiarlos se edita la planilla y se "
          "sube a GitHub, el sitio donde vive la carta. Lo demás pasa solo.", "bajada"),
        pasos([
            ("Descarga la planilla vigente", f"Siempre la última: {enlace(URL_DESCARGA)}"),
            ("Edítala", "En Excel, Numbers o Google Sheets. Cambia precios, agrega productos o sácalos por un tiempo."),
            ("Súbela", f"Arrástrala a {enlace(URL_SUBIDA)} y aprieta el botón verde <b>Commit changes</b>."),
            ("Revisa la carta", f"En unos 2 minutos {enlace(URL_CARTA)} muestra los cambios."),
        ]),
        Spacer(1, 4 * mm),
        recuadro([P("<b>SI ALGO ESTÁ MAL ESCRITO, NO PASA NADA GRAVE</b>", "etiqueta"), Spacer(1, 2),
                  P("Una revisión automática frena la publicación. La carta sigue mostrando la versión anterior, "
                    "los clientes no ven nada roto, y GitHub avisa por correo qué fila corregir.")]),
        Spacer(1, 4 * mm),
        ojo("Si dos personas editan a la vez, queda la planilla de la <b>última</b> que sube. Antes de editar, descarga "
            "siempre la planilla vigente en vez de usar una copia vieja guardada en el computador."),
    ]


def planilla():
    filas_cols = [[P("COLUMNA", "tabla_cab"), P("QUÉ ESCRIBIR", "tabla_cab"), P("EJEMPLO", "tabla_cab")]] + [
        [P(f"<b>{c}</b>", "tabla"), P(q, "tabla"), P(e, "tabla")] for c, q, e in [
            ("Grupo", "El botón de arriba de la carta. Hoy son tres: Combos, Café &amp; bebidas y Para comer.",
             "Café &amp; bebidas"),
            ("Sección", "El subtítulo dentro del grupo.", "Café de especialidad"),
            ("Producto", "El nombre tal como se lee en la carta. Vacío si la fila es una nota.", "Capuccino"),
            ("Detalle", "Opcional: texto chico bajo el nombre. En una nota, el texto de la nota.",
             "César · Quinoa pollo"),
            ("Precio", "Solo el número. Con tamaños: cada tamaño con dos puntos, separados por una barra.",
             "3700<br/>180 ml: 4500 / 500 ml: 9500"),
            ("Mostrar", "<b>Sí</b> para que aparezca en la carta, <b>No</b> para ocultarlo.", "Sí"),
        ]]
    return [
        PageBreak(),
        P("La planilla", "h1"),
        P("La planilla tiene una hoja <b>Carta</b> con una fila por producto y una hoja <b>Instrucciones</b> con este "
          "resumen. La carta respeta el orden de las filas.", "bajada"),
        tabla(filas_cols, [22 * mm, 92 * mm, ANCHO - 114 * mm]),
        Spacer(1, 3 * mm),
        P("Si dejas <b>Grupo</b> o <b>Sección</b> vacíos, se usan los de la fila de arriba. Las filas grises son "
          "<b>notas</b>: no tienen Producto y su texto va en Detalle.", "chico"),
        CondPageBreak(95 * mm),
        P("Así se ve en la planilla…", "h2"),
        muestra_planilla(),
        CondPageBreak(125 * mm),
        P("…y así en la carta", "h2"),
        *figura_carta(),
    ]


def muestra_planilla():
    libro = load_workbook(RAIZ / "datos" / "carta.xlsx", data_only=True, read_only=True)
    filas = [f for f in libro["Carta"].iter_rows(values_only=True)]
    cuerpo = filas[1:]

    def primeras(condicion, n):
        return [f for f in cuerpo if condicion(f)][:n]

    elegidas = (primeras(lambda f: f[1] == "Calientes", 3) + primeras(lambda f: f[1] == "Almuerzos", 3)
                + primeras(lambda f: f[2] == "Helados", 1))

    def precio(v):
        if isinstance(v, (int, float)):
            return "$" + f"{int(v):,}".replace(",", ".")
        return "" if v is None else str(v)

    datos_tabla = [[P(c.upper(), "tabla_cab") for c in filas[0]]]
    notas = []
    for i, f in enumerate(elegidas, start=1):
        es_nota = not f[2]
        estilo = "tabla_nota" if es_nota else "tabla"
        datos_tabla.append([P(str(f[0] or ""), estilo), P(str(f[1] or ""), estilo), P(str(f[2] or ""), estilo),
                            P(str(f[3] or ""), estilo), P(precio(f[4]), estilo), P(str(f[5] or ""), estilo)])
        if es_nota:
            notas.append(i)
    t = tabla(datos_tabla, [18 * mm, 23 * mm, 33 * mm, 42 * mm, ANCHO - 134 * mm, 18 * mm], cebra=False)
    t.setStyle(TableStyle([("BACKGROUND", (0, i), (-1, i), NOTA_FONDO) for i in notas]
                          + [("ALIGN", (4, 1), (4, -1), "RIGHT")]))
    return t


class Maqueta(Flowable):
    """
    Miniatura de la carta dibujada en el propio PDF (no es una captura de pantalla): muestra de dónde sale
    cada columna de la planilla. Al dibujarse acá, no envejece cuando cambia la carta.
    """

    def __init__(self, ancho=ANCHO, alto=98 * mm):
        super().__init__()
        self.width, self.height = ancho, alto

    def _numero(self, n, x, y, d=5 * mm):
        c = self.canv
        c.setFillColor(OLIVA)
        c.circle(x, y, d / 2, stroke=0, fill=1)
        c.setFillColor(CREMA)
        c.setFont("Avenir-Demi", 7.5)
        c.drawCentredString(x, y - 2.5, str(n))

    def _fila(self, x, y, ancho, nombre, precio, tam=8):
        c = self.canv
        c.setFillColor(TINTA)
        c.setFont("Avenir-Demi", tam)
        c.drawString(x, y, nombre)
        ancho_precio = c.stringWidth(precio, "Avenir", tam)
        c.setFont("Avenir", tam)
        c.drawString(x + ancho - ancho_precio, y, precio)
        c.setStrokeColor(LINEA)
        c.setLineWidth(0.5)
        c.setDash(0.5, 2)
        inicio = x + c.stringWidth(nombre, "Avenir-Demi", tam) + 4
        c.line(inicio, y + 1.5, x + ancho - ancho_precio - 4, y + 1.5)
        c.setDash()

    def draw(self):
        c = self.canv
        W, H = self.width, self.height
        margen_num = 8 * mm
        x = margen_num + 5 * mm
        ancho = W - x - 4 * mm
        c.setFillColor(CREMA)
        c.roundRect(margen_num, 0, W - margen_num, H, 3, stroke=0, fill=1)

        y = H - 8 * mm
        # 1 · navegación
        etiquetas = [("Combos", False), ("Café & bebidas", True), ("Para comer", False)]
        cx = x
        for texto, activo in etiquetas:
            an = c.stringWidth(texto, "Avenir", 7) + 8
            c.setFillColor(OLIVA if activo else CREMA)
            c.setStrokeColor(LINEA)
            c.setLineWidth(0.5)
            c.roundRect(cx, y - 2, an, 5.5 * mm / 2 + 4, 4, stroke=1, fill=1)
            c.setFillColor(CREMA if activo else TINTA_SUAVE)
            c.setFont("Avenir", 7)
            c.drawString(cx + 4, y + 1.5, texto)
            cx += an + 4
        self._numero(1, margen_num / 2 + 2 * mm, y + 2)

        # 1 · título del grupo
        y -= 10 * mm
        c.setFillColor(OLIVA)
        c.setFont("Basker", 15)
        c.drawCentredString(x + ancho / 2, y, "Café & bebidas")
        y -= 4 * mm
        c.setStrokeColor(LINEA)
        c.setLineWidth(0.5)
        c.line(x + ancho / 2 - 12, y + 1, x + ancho / 2 - 4, y + 1)
        c.line(x + ancho / 2 + 4, y + 1, x + ancho / 2 + 12, y + 1)
        c.setFillColor(MARCA)
        c.circle(x + ancho / 2, y + 1.5, 0.9, stroke=0, fill=1)

        # 2 · sección
        y -= 8 * mm
        c.setFillColor(MARCA)
        c.setFont("Avenir-Demi", 6.5)
        titulo = "C A F É   D E   E S P E C I A L I D A D"
        c.drawString(x, y, titulo)
        c.setStrokeColor(LINEA)
        c.line(x + c.stringWidth(titulo, "Avenir-Demi", 6.5) + 5, y + 2, x + ancho, y + 2)
        self._numero(2, margen_num / 2 + 2 * mm, y + 1)

        # 3 · nota de sección
        y -= 6 * mm
        c.setFillColor(TINTA_SUAVE)
        c.setFont("Basker-Italica", 8.5)
        c.drawString(x, y, "Trabajamos grano de especialidad, tostado para Boró Café.")
        self._numero(3, margen_num / 2 + 2 * mm, y + 1)

        # 4 · producto con precio · 5 · detalle
        y -= 7 * mm
        self._fila(x, y, ancho, "Espresso", "$3.200")
        self._numero(4, margen_num / 2 + 2 * mm, y + 1)
        y -= 5 * mm
        c.setFillColor(TINTA_SUAVE)
        c.setFont("Basker-Italica", 8)
        c.drawString(x, y, "Corto e intenso, nuestro origen del mes")
        self._numero(5, margen_num / 2 + 2 * mm, y + 1)
        y -= 7 * mm
        self._fila(x, y, ancho, "Cappuccino", "$3.700")

        # 6 · producto con dos precios
        y -= 9 * mm
        c.setFillColor(TINTA)
        c.setFont("Avenir-Demi", 8)
        c.drawString(x, y, "Ensaladas")
        self._numero(6, margen_num / 2 + 2 * mm, y + 1)
        for etiqueta, precio in (("César · Vegetariana", "$6.590"), ("Pollo con quinoa · Salmón", "$6.990")):
            y -= 5.5 * mm
            c.setFillColor(TINTA_SUAVE)
            c.setFont("Basker-Italica", 8)
            c.drawString(x + 3 * mm, y, etiqueta)
            c.setFillColor(TINTA)
            c.setFont("Avenir", 8)
            ap = c.stringWidth(precio, "Avenir", 8)
            c.drawString(x + ancho - ap, y, precio)
            c.setStrokeColor(LINEA)
            c.setDash(0.5, 2)
            c.line(x + 3 * mm + c.stringWidth(etiqueta, "Basker-Italica", 8) + 4, y + 1.5, x + ancho - ap - 4, y + 1.5)
            c.setDash()

        # 7 · promo destacada
        y -= 11 * mm
        alto_promo = 15 * mm
        c.setFillColor(PAPEL_HONDO)
        c.roundRect(x, y - alto_promo + 6 * mm, ancho, alto_promo, 2, stroke=0, fill=1)
        c.setFillColor(TINTA)
        c.setFont("Basker", 11)
        c.drawString(x + 3 * mm, y + 2 * mm, "Tu café")
        c.setFillColor(MARCA)
        c.setFont("Avenir-Demi", 9.5)
        c.drawRightString(x + ancho - 3 * mm, y + 2 * mm, "$2.500")
        c.setFillColor(TINTA_SUAVE)
        c.setFont("Avenir", 7.5)
        c.drawString(x + 3 * mm, y - 2 * mm, "Con cualquier compra de salados o dulces")
        self._numero(7, margen_num / 2 + 2 * mm, y + 2 * mm)


def figura_carta():
    """La miniatura y, debajo, la leyenda de los números en dos columnas."""
    leyenda = [
        ("1", "<b>Grupo</b>: es un botón arriba y el título grande de la sección."),
        ("2", "<b>Sección</b>: el subtítulo verde dentro del grupo."),
        ("3", "<b>Nota</b> de sección: fila sin Producto, con el texto en Detalle."),
        ("4", "<b>Producto</b> y <b>Precio</b>."),
        ("5", "<b>Detalle</b>: el texto chico bajo el nombre."),
        ("6", "Un producto con <b>dos precios</b>: en Precio va «César · Vegetariana: 6590 / Pollo con "
              "quinoa · Salmón: 6990»."),
        ("7", "Sección cuyo nombre empieza con <b>«Promo»</b>: sus productos salen destacados en recuadros."),
    ]
    sin_relleno = [("VALIGN", (0, 0), (-1, -1), "TOP"),
                   ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                   ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]
    media = ANCHO / 2

    def item_leyenda(n, texto):
        c = Table([[Numero(n, 5.5 * mm), P(texto, "chico")]], colWidths=[8 * mm, media - 8 * mm - 6 * mm])
        c.setStyle(TableStyle(sin_relleno))
        return c

    filas = []
    for i in range(4):
        izq = item_leyenda(*leyenda[i])
        der = item_leyenda(*leyenda[i + 4]) if i + 4 < len(leyenda) else ""
        filas.append([izq, der])
    leyenda_tabla = Table(filas, colWidths=[media, media])
    leyenda_tabla.setStyle(TableStyle(sin_relleno + [("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    return [Maqueta(), Spacer(1, 5 * mm), leyenda_tabla]


def tareas():
    filas = [[P("QUIERO…", "tabla_cab"), P("CÓMO SE HACE", "tabla_cab")]] + [
        [P(f"<b>{q}</b>", "tabla"), P(c, "tabla")] for q, c in [
            ("Cambiar un precio",
             "Escribe el precio nuevo en <b>Precio</b>. Solo el número: <b>3900</b> (también se entiende $3.900)."),
            ("Agregar un producto",
             "Inserta una fila debajo de otro producto de la misma sección (clic derecho en el número de fila › "
             "Insertar). Completa <b>Producto</b>, <b>Precio</b> y <b>Mostrar: Sí</b>. Grupo y Sección pueden "
             "quedar vacíos: se toman de la fila de arriba."),
            ("Sacar un producto por un tiempo",
             "Agotado o de temporada: cambia <b>Mostrar</b> a <b>No</b>. La fila queda guardada; para que vuelva, "
             "pon <b>Sí</b>."),
            ("Eliminar un producto para siempre", "Borra la fila completa (clic derecho en el número de fila › Eliminar)."),
            ("Un producto con dos precios o tamaños",
             "En <b>Precio</b>: <b>180 ml: 4500 / 500 ml: 9500</b>, o <b>César · Vegetariana: 6590 / Pollo con "
             "quinoa: 6990</b>. Cada uno sale en su propia línea."),
            ("Un precio único para toda una sección",
             "Al comienzo de la sección, agrega una fila <b>sin Producto</b> con <b>Todos a $6.990</b> en "
             "<b>Detalle</b>. Los productos de esa sección pueden ir sin precio."),
            ("Una nota bajo el título de la sección",
             "Igual que la anterior, con cualquier texto en Detalle. Por ejemplo: <i>Leche vegetal en cualquier café "
             "+$800</i>."),
            ("Una sección o un grupo nuevo",
             "Escribe un nombre nuevo en <b>Sección</b> o en <b>Grupo</b>. Un grupo nuevo aparece como un botón más "
             "arriba de la carta, sin ilustración."),
            ("Destacar una promoción",
             "Ponla en una sección cuyo nombre empiece con <b>Promo</b> (por ejemplo «Promos de la mañana»): "
             "sus productos se muestran en recuadros, no como lista de precios."),
            ("Cambiar el orden", "La carta sigue el orden de las filas: corta la fila y pégala donde la quieras."),
        ]]
    return [
        PageBreak(),
        P("Tareas comunes", "h1"),
        P("Todo se hace en la hoja <b>Carta</b>. Después de cualquier cambio, guarda y súbela (página siguiente).",
          "bajada"),
        tabla(filas, [48 * mm, ANCHO - 48 * mm]),
    ]


def publicar():
    return [
        PageBreak(),
        P("Publicar los cambios", "h1"),
        P("Para subir la planilla necesitas una cuenta de GitHub con acceso a la organización <b>borocafe</b> "
          "(ver Preguntas frecuentes).", "bajada"),
        pasos([
            ("Guarda la planilla",
             f"En Excel o Numbers: {ruta_menu('Archivo', 'Guardar')}, en formato .xlsx. "
             f"En Google Sheets: {ruta_menu('Archivo', 'Descargar', 'Microsoft Excel (.xlsx)')}."),
            ("Abre la página de subida",
             f"{enlace(URL_SUBIDA)}. Si te lo pide, inicia sesión con tu cuenta de GitHub."),
            ("Arrastra el archivo",
             "Suéltalo en el recuadro de la página, o toca <b>choose your files</b> y elígelo. El nombre del archivo "
             "da igual."),
            ("Anota qué cambiaste (opcional)",
             "En el cuadro <b>Commit changes</b>, por ejemplo: <i>Capuccino a $3.900</i>. Queda en el historial."),
            ("Aprieta el botón verde «Commit changes»", "Listo: la revisión automática empieza sola."),
            ("Revisa la carta",
             f"En unos 2 minutos abre {enlace(URL_CARTA)}. Si ves el precio anterior, recarga la página."),
        ]),
        P("También se puede hacer desde el navegador del celular.", "chico"),
        P("¿Cómo sé si se publicó?", "h2"),
        P(f"Entra a {enlace(URL_ACTIONS)}. La primera fila de la lista es tu subida:"),
        Spacer(1, 2 * mm),
        tabla([
            [P("<b>Círculo amarillo</b>", "tabla"), P("Se está revisando y publicando. Espera un momento.", "tabla")],
            [P("<b>Tic verde</b>", "tabla"), P("Publicado. La carta ya tiene tus cambios.", "tabla")],
            [P("<b>Cruz roja</b>", "tabla"),
             P("No se publicó porque la planilla tiene algo que corregir. Ve a <b>Si algo sale mal</b>.", "tabla")],
        ], [38 * mm, ANCHO - 38 * mm], cabecera=False),
    ]


def errores():
    filas = [[P("EL AVISO DICE…", "tabla_cab"), P("QUÉ HACER", "tabla_cab")]] + [
        [P(a, "tabla"), P(b, "tabla")] for a, b in [
            ("Fila 12 (Latte): el precio «tres mil» no se entiende",
             "Escribe solo números en Precio (3700), o tamaños así: 180 ml: 4500 / 500 ml: 9500."),
            ("Fila 8 (Cortado): falta el precio",
             "Pon el precio. Si toda la sección cuesta lo mismo, agrega una nota «Todos a $...»."),
            ("Fila 5 (Moka): en Mostrar escribe Sí o No", "Corrige esa celda: solo Sí o No."),
            ("Fila 2: falta el Grupo", "La primera fila de la planilla tiene que decir su Grupo (Combos, "
             "Café &amp; bebidas o Para comer)."),
            ("Faltan las columnas… / no encontré la fila de títulos",
             "Se cambió o borró la primera fila. Descarga de nuevo la planilla vigente y repite tus cambios."),
            ("La carta quedaría con solo 3 productos visibles",
             "Probablemente se subió otro archivo, o muchas filas quedaron en No. Revisa y vuelve a subir."),
        ]]
    return [
        PageBreak(),
        P("Si algo sale mal", "h1"),
        P("Si la revisión encuentra un problema, <b>la carta no cambia</b>: sigue la versión anterior y los clientes no "
          "notan nada.", "bajada"),
        pasos([
            ("Lee el aviso",
             "Llega un correo de GitHub (<i>Run failed: Publicar carta</i>). También lo ves en "
             f"{enlace(URL_ACTIONS)}: toca la fila con la cruz roja y busca el aviso <b>Revisa la planilla</b>."),
            ("Corrige la fila que indica", "El aviso dice el número de fila y el producto."),
            ("Vuelve a subir la planilla", "Igual que siempre. Si todo está bien, se publica."),
        ]),
        Spacer(1, 2 * mm),
        tabla(filas, [72 * mm, ANCHO - 72 * mm]),
        CondPageBreak(60 * mm),
        P("Volver a una versión anterior", "h2"),
        P("Cada subida queda guardada. Si un cambio salió mal, puedes recuperar la planilla de antes:"),
        Spacer(1, 2 * mm),
        pasos([
            ("", f"Abre {enlace(URL_HISTORIAL)}: la lista de subidas, la más nueva arriba."),
            ("", "En la versión que quieres recuperar, toca el botón <b>&lt;&gt;</b> (<i>Browse repository at this "
                 "point</i>)."),
            ("", f"Entra a {ruta_menu('datos', 'carta.xlsx')} y toca <b>Download raw file</b> (ícono de descarga)."),
            ("", "Sube ese archivo como siempre."),
        ]),
    ]


def qr_nfc():
    return [
        PageBreak(),
        P("Tarjetas de mesa: QR y NFC", "h1"),
        P(f"El QR y el sticker NFC llevan a la misma dirección, {enlace(URL_QR)}, que abre la carta. Si algún día la "
          "carta cambia de dirección, se corrige esa redirección y las tarjetas siguen sirviendo.", "bajada"),
        P("Imprimir la tarjeta", "h2"),
        P(f"Descarga {enlace(URL_TARJETA, 'la tarjeta en PDF')} e imprímela al <b>100 %</b>, sin «ajustar a la "
          "página»: mide 10,5 × 14,8 cm (A6). Mejor en cartulina. <b>Antes de imprimir muchas, escanea una con el "
          "celular.</b>"),
        P("Programar el sticker NFC", "h2"),
        pasos([
            ("Compra los stickers",
             "NFC <b>NTAG213</b> o <b>NTAG215</b>, redondos de 25 a 30 mm. Para mesas o bandejas de metal, "
             "de tipo <b>anti-metal</b>."),
            ("Instala NFC Tools", "App gratuita para iPhone y Android."),
            ("Graba la dirección",
             f"En NFC Tools: {ruta_menu('Escribir', 'Agregar un registro', 'URL')}. Escribe "
             "<b>https://borocafe.github.io/carta/ir/</b>, toca <b>Escribir</b> y acerca el sticker a la parte de "
             "arriba del iPhone o a la parte de atrás del Android."),
            ("Pruébalo", "Con la pantalla desbloqueada, acerca el celular: tiene que abrir la carta."),
            ("Bloquéalo",
             f"{ruta_menu('Otros', 'Bloquear etiqueta')}. Así nadie puede cambiar la dirección, ni siquiera tú, y no "
             "hace falta: la redirección se ajusta sin tocar el sticker."),
            ("Pégalo detrás del QR", "Idealmente entre dos capas o bajo el plastificado, para que no se vea."),
        ]),
        Spacer(1, 2 * mm),
        P("Los nombres de los botones pueden variar un poco según la versión de la app.", "chico"),
        Spacer(1, 4 * mm),
        ojo("Una vez por semana revisa que nadie haya pegado otro QR o sticker encima de las tarjetas. Si ves algo "
            "raro, cambia la tarjeta: alguien podría intentar llevar a los clientes a otra página."),
    ]


def preguntas():
    lista = [
        ("¿Puedo cambiar el diseño, las ilustraciones, la dirección o el horario desde la planilla?",
         "No. La planilla es solo para productos, precios y notas. Esos cambios los hace quien administra la carta."),
        ("¿Quién puede subir cambios?",
         "Cualquier persona con una cuenta gratuita de GitHub invitada a la organización. Se invita en "
         f"{enlace('https://github.com/orgs/borocafe/people', 'github.com/orgs/borocafe/people')} › "
         "<b>Invite member</b>, con permiso de escritura en el repositorio <b>carta</b>."),
        ("¿Cuánto demora en verse un cambio?", "Unos 2 minutos desde que aprietas Commit changes."),
        ("¿Importa el nombre del archivo que subo?",
         "No. Se usa la última planilla subida y queda guardada como carta.xlsx. Sirve .xlsx o .csv."),
        ("¿Se puede deshacer un cambio?", "Sí: ver <b>Volver a una versión anterior</b>."),
        ("¿Qué pasa si falla la publicación automática?",
         "La carta publicada sigue igual. Revisa el aviso y, si no hay nada que corregir, vuelve a subir la planilla "
         "más tarde."),
    ]
    contenido = [PageBreak(), P("Preguntas frecuentes", "h1"), Spacer(1, 2 * mm)]
    for pregunta, respuesta in lista:
        contenido.append(KeepTogether([P(f"<b>{pregunta}</b>"), Spacer(1, 1.5 * mm), P(respuesta), Spacer(1, 5 * mm)]))
    return contenido


# ---------------------------------------------------------------- páginas

def fondo_portada(canv, doc):
    canv.saveState()
    canv.setFillColor(CREMA)
    canv.rect(0, 0, ANCHO_PAGINA, ALTO_PAGINA, stroke=0, fill=1)
    canv.restoreState()


def encabezado(canv, doc):
    canv.saveState()
    canv.setFont("Avenir", 8)
    canv.setFillColor(TINTA_SUAVE)
    y = ALTO_PAGINA - 14 * mm
    canv.drawString(MARGEN_X, y, "Boró Café · Manual de la carta")
    canv.drawRightString(ANCHO_PAGINA - MARGEN_X, y, str(doc.page))
    canv.setStrokeColor(LINEA)
    canv.setLineWidth(0.5)
    canv.line(MARGEN_X, y - 3 * mm, ANCHO_PAGINA - MARGEN_X, y - 3 * mm)
    canv.restoreState()


def revisar_caracteres():
    faltan = {}
    for texto, fuente in TEXTOS:
        plano = re.sub(r"<[^>]+>", "", texto).replace("&lt;", "<").replace("&gt;", ">")
        cara = pdfmetrics.getFont(fuente).face
        for ch in set(plano):
            if ord(ch) > 32 and ord(ch) not in cara.charToGlyph:
                faltan.setdefault(fuente, set()).add(ch)
    return faltan


def main():
    registrar_fuentes()
    estilos()
    MANUAL.mkdir(exist_ok=True)
    doc = BaseDocTemplate(str(SALIDA), pagesize=A4, title="Manual de la carta · Boró Café", author="Boró Café",
                          subject="Cómo actualizar la carta digital", leftMargin=MARGEN_X, rightMargin=MARGEN_X,
                          topMargin=MARGEN_SUP, bottomMargin=MARGEN_INF)
    marco = Frame(MARGEN_X, MARGEN_INF, ANCHO, ALTO_PAGINA - MARGEN_SUP - MARGEN_INF, id="marco",
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc.addPageTemplates([PageTemplate("portada", [marco], onPage=fondo_portada),
                          PageTemplate("interior", [marco], onPage=encabezado)])
    historia = portada() + resumen() + planilla() + tareas() + publicar() + errores() + qr_nfc() + preguntas()
    faltan = revisar_caracteres()
    if faltan:
        sys.exit(f"Caracteres sin glifo en las fuentes: {faltan}")
    doc.build(historia)
    print(f"escrito: {SALIDA.relative_to(RAIZ)} ({SALIDA.stat().st_size // 1024} KB, {doc.page} páginas)")


if __name__ == "__main__":
    main()
