#!/usr/bin/env python3
"""
Escribe la planilla Excel de la carta (datos/carta.xlsx) con formato, listas desplegables e instrucciones.

Uso:
    python3 herramientas/planilla.py desde-json <carta-datos.json>   # primera versión, desde la carta extraída
    python3 herramientas/planilla.py normalizar                      # lo corre GitHub Actions tras cada subida

`normalizar` toma la planilla subida más recientemente (cualquier nombre, .xlsx o .csv), la reescribe como
datos/carta.xlsx y borra las demás. Así siempre hay un solo archivo para descargar y editar.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import datos  # noqa: E402

OLIVA, CREMA, NOTA_FONDO, TINTA_SUAVE = "4B5324", "F4E8D0", "F6F0E2", "6D5B48"
ANCHOS = {"Grupo": 12, "Sección": 24, "Producto": 42, "Detalle": 50, "Precio": 30, "Mostrar": 10}
DESTINO = datos.DATOS / "carta.xlsx"

INSTRUCCIONES = [
    ("Carta de Boró Café: cómo actualizarla", "titulo"),
    ("", None),
    ("EDITAR (hoja «Carta»)", "sub"),
    ("• Cambiar un precio: escribe el número nuevo en Precio. Solo el número, por ejemplo 3700.", None),
    ("• Agregar un producto: inserta una fila dentro de su sección y completa Producto y Precio. "
     "Si dejas Grupo y Sección vacíos, se usan los de la fila de arriba.", None),
    ("• Sacar un producto por un tiempo (agotado, de temporada): pon No en Mostrar. Para que vuelva, Sí.", None),
    ("• Detalle (opcional): texto chico bajo el nombre, por ejemplo «César · Quinoa pollo · Vegetariana».", None),
    ("• Productos con tamaños: en Precio escribe «180 ml: 4500 / 500 ml: 9500».", None),
    ("• Notas de sección (filas grises): van SIN Producto, con el texto en Detalle. Si la nota empieza con "
     "«Todos a» (por ejemplo «Todos a $6.990»), se muestra como sello y los productos de esa sección van sin precio.", None),
    ("• La carta respeta el orden de las filas. Para una sección o un grupo nuevo, escribe un nombre nuevo: "
     "cada grupo es un botón arriba de la carta.", None),
    ("", None),
    ("PUBLICAR", "sub"),
    ("1. Guarda el archivo (sirve Excel, Numbers o Google Sheets descargado como .xlsx o .csv).", None),
    ("2. Entra a https://github.com/borocafe/carta/upload/main/datos", None),
    ("3. Arrastra el archivo (puede tener cualquier nombre) y aprieta «Commit changes».", None),
    ("4. En unos 2 minutos la carta está actualizada: https://borocafe.github.io/carta/", None),
    ("", None),
    ("Si algo está mal escrito, la carta NO cambia y GitHub avisa por correo qué fila revisar. "
     "Corrígela y vuelve a subir el archivo.", None),
    ("Para editar siempre la versión más reciente, descárgala de: https://github.com/borocafe/carta/raw/main/datos/carta.xlsx", None),
]


def texto_precio(p):
    if isinstance(p, list):
        return " / ".join(f"{e}: {v}" for e, v in p)
    return p


def escribir(filas, destino=DESTINO):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Carta"
    hoja.append(datos.COLUMNAS)
    for c in hoja[1]:
        c.font = Font(bold=True, color=CREMA)
        c.fill = PatternFill("solid", fgColor=OLIVA)
        c.alignment = Alignment(vertical="center")
    hoja.row_dimensions[1].height = 24

    linea = Side(style="thin", color="D9C7A7")
    anterior = None
    for f in filas:
        valor = texto_precio(f["precio"])
        hoja.append([f["grupo"], f["seccion"], f["producto"] or None, f["detalle"] or None, valor,
                     "Sí" if f["mostrar"] else "No"])
        r = hoja.max_row
        precio = hoja.cell(r, 5)
        precio.alignment = Alignment(horizontal="right")
        if isinstance(valor, int):
            precio.number_format = '"$"#,##0'
        if f["tipo"] == "nota":
            for col in range(1, 7):
                hoja.cell(r, col).fill = PatternFill("solid", fgColor=NOTA_FONDO)
                hoja.cell(r, col).font = Font(italic=True, color=TINTA_SUAVE)
        clave = (datos.normal(f["grupo"]), datos.normal(f["seccion"]))
        if anterior is not None and clave != anterior:
            for col in range(1, 7):
                hoja.cell(r, col).border = Border(top=linea)
        anterior = clave

    for i, col in enumerate(datos.COLUMNAS, start=1):
        hoja.column_dimensions[chr(64 + i)].width = ANCHOS[col]
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = f"A1:F{hoja.max_row}"

    ultima = hoja.max_row + 300
    si_no = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True, showErrorMessage=True,
                           errorTitle="Mostrar", error="Escribe Sí o No.")
    si_no.add(f"F2:F{ultima}")
    nombres = list(dict.fromkeys(f["grupo"] for f in filas))
    grupo = DataValidation(type="list", formula1='"' + ",".join(nombres) + '"', allow_blank=True,
                           showErrorMessage=True, errorStyle="information", errorTitle="Grupo nuevo",
                           error="Ese grupo todavía no existe: la carta va a mostrar un botón nuevo.")
    grupo.add(f"A2:A{ultima}")
    hoja.add_data_validation(si_no)
    hoja.add_data_validation(grupo)

    ayuda = libro.create_sheet("Instrucciones")
    ayuda.column_dimensions["A"].width = 110
    for texto, estilo in INSTRUCCIONES:
        ayuda.append([texto])
        c = ayuda.cell(ayuda.max_row, 1)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if estilo == "titulo":
            c.font = Font(bold=True, size=15, color=OLIVA)
        elif estilo == "sub":
            c.font = Font(bold=True, color=OLIVA)

    libro.properties.title = "Carta Boró Café"
    destino.parent.mkdir(parents=True, exist_ok=True)
    libro.save(destino)


def desde_json(ruta):
    d = json.loads(Path(ruta).read_text(encoding="utf-8"))
    filas = []
    for s in d["secciones"]:
        for texto in (s["precio_fijo"], s["nota"]):
            if texto:
                filas.append(dict(tipo="nota", grupo=s["grupo"], seccion=s["seccion"], producto="", detalle=texto,
                                  precio=None, mostrar=True))
        for p in d["productos"]:
            if (p["grupo"], p["seccion"]) == (s["grupo"], s["seccion"]):
                filas.append(dict(tipo="producto", grupo=p["grupo"], seccion=p["seccion"], producto=p["producto"],
                                  detalle=p["detalle"], precio=datos.precio(p["precio"]), mostrar=True))
    escribir(filas)
    print(f"escrito: datos/carta.xlsx ({len(filas)} filas)")


def normalizar():
    vigente = datos.archivo_vigente()
    if vigente == DESTINO and len(datos.planillas()) == 1:
        print("datos/carta.xlsx ya es la planilla vigente")
        return
    _, filas = datos.validar(datos.leer(vigente), vigente.name)
    escribir(filas)
    for p in datos.planillas():
        if p != DESTINO:
            p.unlink()
    print(f"datos/{vigente.name} → datos/carta.xlsx")


if __name__ == "__main__":
    if len(sys.argv) >= 3 and sys.argv[1] == "desde-json":
        desde_json(sys.argv[2])
    elif len(sys.argv) == 2 and sys.argv[1] == "normalizar":
        normalizar()
    else:
        sys.exit(__doc__)
