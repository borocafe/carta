#!/usr/bin/env python3
"""
Lectura y validación de la planilla de la carta (datos/*.xlsx o datos/*.csv).

Reglas de la planilla (las mismas que explica la hoja "Instrucciones" del Excel):
  - Columnas: Grupo, Sección, Producto, Detalle, Precio, Mostrar. El orden de las filas es el de la carta.
  - Grupo o Sección vacíos heredan el valor de la fila de arriba.
  - Fila sin Producto y con Detalle = nota de la sección. Si empieza con "Todos a" es un precio fijo,
    y los productos de esa sección pueden ir sin precio.
  - Precio: número (3700, $3.700, 3.700) o tamaños: "180 ml: 4500 / 500 ml: 9500".
  - Mostrar: Sí / No (vacío cuenta como Sí).
"""
import csv
import io
import re
import subprocess
import unicodedata
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
DATOS = RAIZ / "datos"
COLUMNAS = ["Grupo", "Sección", "Producto", "Detalle", "Precio", "Mostrar"]
SI = {"", "si", "s", "x", "true", "verdadero", "1"}
NO = {"no", "n", "false", "falso", "0"}
MINIMO_PRODUCTOS = 5


class ErrorPlanilla(Exception):
    def __init__(self, errores):
        super().__init__("\n".join(errores))
        self.errores = errores


def normal(texto):
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().lower()


def slug(texto):
    return re.sub(r"[^a-z0-9]+", "-", normal(texto)).strip("-") or "seccion"


def planillas():
    return [p for p in DATOS.glob("*") if p.suffix.lower() in (".xlsx", ".csv") and not p.name.startswith(("~$", "."))]


def archivo_vigente():
    """La planilla subida más recientemente a datos/, por fecha de commit. Un archivo sin commitear gana."""
    candidatos = planillas()
    if not candidatos:
        raise ErrorPlanilla(["No hay ninguna planilla en la carpeta datos/ (se espera un archivo .xlsx o .csv)."])

    def fecha(p):
        r = subprocess.run(["git", "log", "-1", "--format=%ct", "--", p.name], cwd=DATOS, capture_output=True, text=True)
        commit = r.stdout.strip()
        return (float(commit) if commit else float("inf"), p.stat().st_mtime)

    return max(candidatos, key=fecha)


def leer(ruta):
    """La tabla como lista de filas (listas de celdas)."""
    ruta = Path(ruta)
    if ruta.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        libro = load_workbook(ruta, data_only=True, read_only=True)
        hoja = next((h for h in libro.worksheets if normal(h.title) == "carta"), libro.worksheets[0])
        return [["" if c is None else c for c in fila] for fila in hoja.iter_rows(values_only=True)]
    crudo = ruta.read_bytes()
    try:
        texto = crudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = crudo.decode("cp1252")  # Excel en Windows con configuración en español
    primera = texto.splitlines()[0] if texto.strip() else ""
    separador = ";" if primera.count(";") > primera.count(",") else ","
    return list(csv.reader(io.StringIO(texto), delimiter=separador))


def _numero(texto):
    t = str(texto).replace("$", "").replace(" ", "").replace(" ", "").strip()
    if re.fullmatch(r"\d{1,3}([.,]\d{3})+", t):  # 4.500 / 4,500
        return int(re.sub(r"[.,]", "", t))
    m = re.fullmatch(r"(\d+)([.,]0{1,2})?", t)  # 4500 / 4500,00
    return int(m.group(1)) if m else None


def precio(valor):
    """None si está vacío; un entero; o [(etiqueta, entero)] para tamaños. ValueError si no se entiende."""
    if valor is None or str(valor).strip() == "":
        return None
    if isinstance(valor, bool):
        raise ValueError(valor)
    if isinstance(valor, (int, float)):
        if valor < 0 or valor != int(valor):
            raise ValueError(valor)
        return int(valor)
    texto = str(valor).strip()
    if ":" in texto:
        tamanos = []
        for parte in re.split(r"\s*/\s*", texto):
            etiqueta, _, monto = parte.partition(":")
            n = _numero(monto)
            if not etiqueta.strip() or n is None:
                raise ValueError(valor)
            tamanos.append((etiqueta.strip(), n))
        return tamanos
    n = _numero(texto)
    if n is None:
        raise ValueError(valor)
    return n


def validar(tabla, nombre_archivo="La planilla"):
    """Devuelve (grupos visibles, todas las filas). Lanza ErrorPlanilla con mensajes para personas."""
    idx_cab, mapa = None, {}
    for i, fila in enumerate(tabla[:10]):
        nombres = [normal(c) for c in fila]
        if "producto" in nombres and "precio" in nombres:
            idx_cab = i
            mapa = {col: nombres.index(normal(col)) for col in COLUMNAS if normal(col) in nombres}
            break
    if idx_cab is None:
        raise ErrorPlanilla([f"{nombre_archivo}: no encontré la fila de títulos (Grupo, Sección, Producto, Detalle, Precio, Mostrar)."])
    faltan = [c for c in COLUMNAS if c not in mapa]
    if faltan:
        raise ErrorPlanilla([f"{nombre_archivo}: faltan las columnas {', '.join(faltan)}."])

    def celda(fila, col):
        i = mapa[col]
        return fila[i] if i < len(fila) else ""

    errores, filas, grupo, seccion = [], [], "", ""
    for n, fila in enumerate(tabla[idx_cab + 1:], start=idx_cab + 2):
        if all(str(c).strip() == "" for c in fila):
            continue
        grupo = str(celda(fila, "Grupo")).strip() or grupo
        seccion = str(celda(fila, "Sección")).strip() or seccion
        producto = str(celda(fila, "Producto")).strip()
        detalle = str(celda(fila, "Detalle")).strip()
        donde = f"Fila {n}" + (f" ({producto})" if producto else "")
        if not producto and not detalle:
            continue
        if not grupo:
            errores.append(f"{donde}: falta el Grupo (por ejemplo Café, Comida, Dulces o Pan).")
            continue
        if not seccion:
            errores.append(f"{donde}: falta la Sección.")
            continue
        mostrar = normal(celda(fila, "Mostrar"))
        if mostrar not in SI | NO:
            errores.append(f"{donde}: en Mostrar escribe Sí o No (dice «{celda(fila, 'Mostrar')}»).")
            continue
        base = dict(grupo=grupo, seccion=seccion, detalle=detalle, mostrar=mostrar not in NO, fila=n)
        if not producto:
            filas.append(dict(base, tipo="nota", producto="", precio=None))
            continue
        try:
            p = precio(celda(fila, "Precio"))
        except ValueError:
            errores.append(f"{donde}: el precio «{celda(fila, 'Precio')}» no se entiende. Escribe solo el número "
                           f"(por ejemplo 3700) o tamaños así: 180 ml: 4500 / 500 ml: 9500.")
            continue
        filas.append(dict(base, tipo="producto", producto=producto, precio=p))

    grupos = []
    for f in filas:
        if not f["mostrar"]:
            continue
        g = next((x for x in grupos if normal(x["nombre"]) == normal(f["grupo"])), None)
        if g is None:
            g = dict(nombre=f["grupo"], secciones=[])
            grupos.append(g)
        s = next((x for x in g["secciones"] if normal(x["nombre"]) == normal(f["seccion"])), None)
        if s is None:
            s = dict(nombre=f["seccion"], notas=[], productos=[])
            g["secciones"].append(s)
        (s["notas"] if f["tipo"] == "nota" else s["productos"]).append(f)

    for g in grupos:
        for s in g["secciones"]:
            fijo = any(normal(nt["detalle"]).startswith("todos a") for nt in s["notas"])
            for p in s["productos"]:
                if p["precio"] is None and not fijo:
                    errores.append(f"Fila {p['fila']} ({p['producto']}): falta el precio. Si toda la sección cuesta "
                                   f"lo mismo, agrega una nota «Todos a $...».")
        g["secciones"] = [s for s in g["secciones"] if s["productos"]]
    grupos = [g for g in grupos if g["secciones"]]

    total = sum(len(s["productos"]) for g in grupos for s in g["secciones"])
    if not errores and total < MINIMO_PRODUCTOS:
        errores.append(f"{nombre_archivo}: la carta quedaría con solo {total} productos visibles. ¿Se subió el archivo correcto?")
    if errores:
        raise ErrorPlanilla(errores)
    return grupos, filas
